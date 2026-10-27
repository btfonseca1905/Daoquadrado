#!/usr/bin/env python3
"""Atualiza o catálogo Markdown a partir das cinco abas da base original.

Uso (opcional): python scripts/atualizar_catalogo.py
Requer openpyxl somente para atualizar a base; a Skill funciona sem executá-lo.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit("Para atualizar o catálogo, instale a dependência: pip install openpyxl") from exc


COLUNAS = {
    "Métricas de Negócio": {"nome": 3, "sigla": 4, "finalidade": 5, "calculo": 6, "extras": [("Tipo de métrica", 2)]},
    "Métricas de Marketing": {"nome": 4, "sigla": 5, "finalidade": 6, "calculo": 7, "extras": [("Objetivo de uso", 2), ("Etapa do funil", 3)]},
    "Métricas de Produto": {"nome": 5, "sigla": 6, "finalidade": 7, "calculo": 8, "extras": [("Foco de trabalho", 2), ("Etapa do funil", 3), ("Estágio", 4)]},
    "Métricas de UX": {"nome": 3, "sigla": 4, "finalidade": 5, "calculo": 6, "extras": [("Grupo de métricas", 2)]},
    "Métricas de Delivery": {"nome": 3, "sigla": 4, "finalidade": 5, "calculo": 6, "extras": [("Grupo de métricas", 2)]},
}


def inline(value: object) -> str:
    if value is None:
        return "Não informado"
    value = re.sub(r"\s+", " ", str(value).strip())
    return (value or "Não informado").replace("|", r"\|")


def build_catalog(workbook: Path) -> tuple[str, dict[str, int]]:
    book = openpyxl.load_workbook(workbook, data_only=False)
    if set(book.sheetnames) != set(COLUNAS):
        raise ValueError(f"Abas esperadas: {list(COLUNAS)}. Abas recebidas: {book.sheetnames}")
    output = [
        "# Catálogo de métricas de Product Management",
        "",
        "Fonte: `assets/base-metricas.xlsx`, planilha fornecida por Bruna Fonseca.",
        "",
        "Este índice reproduz a descrição do campo **Como medir** da base, sem corrigir fórmulas; por isso, leia também `cuidados-formulas.md`.",
        "O sufixo `!n` indica o número da linha no Excel. Uma mesma métrica pode constar em mais de uma aba.",
        "",
    ]
    counts: dict[str, int] = {}
    for title, config in COLUNAS.items():
        sheet = book[title]
        output.extend([f"## {title}", ""])
        count = 0
        for row in range(2, sheet.max_row + 1):
            name = sheet.cell(row, config["nome"]).value
            if not name:
                continue
            count += 1
            output.extend([
                f"### {inline(name)}",
                "",
                f"- **Origem:** `{title}!{row}`",
                f"- **Sigla / nome simplificado:** {inline(sheet.cell(row, config['sigla']).value)}",
            ])
            for label, col in config["extras"]:
                output.append(f"- **{label}:** {inline(sheet.cell(row, col).value)}")
            output.extend([
                f"- **Função da métrica:** {inline(sheet.cell(row, config['finalidade']).value)}",
                f"- **Como medir (texto da planilha):** {inline(sheet.cell(row, config['calculo']).value)}",
                "",
            ])
        counts[title] = count
    output[7:7] = ["## Quantidade de registros", "", *(f"- {sheet}: {count}" for sheet, count in counts.items()), f"- **Total: {sum(counts.values())}**", ""]
    return "\n".join(output) + "\n", counts


def main() -> None:
    skill_folder = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--planilha", type=Path, default=skill_folder / "assets" / "base-metricas.xlsx")
    parser.add_argument("--saida", type=Path, default=skill_folder / "references" / "catalogo-metricas.md")
    args = parser.parse_args()
    catalog, counts = build_catalog(args.planilha)
    args.saida.write_text(catalog, encoding="utf-8")
    print(f"Catálogo atualizado: {sum(counts.values())} métricas nas {len(counts)} abas, {args.saida}")


if __name__ == "__main__":
    main()
